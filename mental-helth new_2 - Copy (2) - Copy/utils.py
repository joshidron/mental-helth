import os
from docx import Document
from fpdf import FPDF
from datetime import datetime
import re

# Font paths
FONT_DIR = "static/fonts"
FONT_PATH_NIRMALA = os.path.join(FONT_DIR, "Nirmala.ttc")

# Localization Dictionary for PDF
PDF_TRANSLATIONS = {
    'en': {
        'title': 'Prajna Path Report',
        'subtitle': 'Your Personalized Prajna Path Journey',
        'name': 'Name',
        'age': 'Age',
        'profession': 'Profession',
        'date': 'Date',
        'disclaimer': 'Disclaimer: This report is for informational purposes only and does not replace professional medical advice.',
        'page': 'Page',
        'no_data': 'No specific advice found for this section. Please consult a specialist.',
        'main symptoms': 'Main Symptoms',
        'myths': 'Myths',
        'routine': 'Daily Routine',
        'social': 'Social Activities',
        'food': 'Food Habits'
    },
    'hi': {
        'title': 'Prajna Path Report',
        'subtitle': 'आपकी व्यक्तिगत मानसिक स्वास्थ्य यात्रा',
        'name': 'नाम',
        'age': 'उम्र',
        'profession': 'पेशा',
        'date': 'तारीख',
        'disclaimer': 'Disclaimer: This report is for informational purposes only and does not replace professional medical advice.',
        'page': 'Page',
        'no_data': 'कृपया किसी विशेषज्ञ से सलाह लें।',
        'gad': 'सामान्यीकृत चिंता विकार (GAD)',
        'depression': 'अवसाद (Depression)',
        'stress': 'तनाव और बर्नआउट (Stress & Burnout)',
        'ocd': 'जुनून-बाध्यकारी विकार (OCD)',
        'ptsd': 'पोस्ट-ट्रॉमैटिक स्ट्रेस (PTSD)',
        'personality': 'व्यक्तित्व असंतुलन (Personality Imbalance)',
        'adhd': 'एडीएचडी (ADHD)',
        'autism': 'ऑटिज़्म स्पेक्ट्रम (Autism)',
        'eating': 'खाने के विकार (Eating Disorders)',
        'main symptoms': 'मुख्य लक्षण',
        'myths': 'भ्रामक धारणाएं (मिथक)',
        'routine': 'दैनिक दिनचर्या',
        'social': 'सामाजिक गतिविधियाँ',
        'food': 'भोजन की आदतें'
    },
    'gu': {
        'title': 'Prajna Path Report',
        'subtitle': 'તમારી વ્યક્તિગત માનસિક સ્વાસ્થ્ય યાત્રા',
        'name': 'નામ',
        'age': 'ઉંમર',
        'profession': 'વ્યવસાય',
        'date': 'તારીખ',
        'disclaimer': 'Disclaimer: This report is for informational purposes only and does not replace professional medical advice.',
        'page': 'Page',
        'no_data': 'આ વિભાગ માટે કોઈ ચોક્કસ સલાહ મળી નથી. કૃપા કરી નિષ્ણાતની સલાહ લો.',
        'gad': 'ચિંતાનો રોગ (GAD)',
        'depression': 'ડિપ્રેશન',
        'stress': 'તણાવ અને થાક',
        'ocd': 'ઓસીડી (OCD)',
        'ptsd': 'પીટીએસડી (PTSD)',
        'personality': 'વ્યક્તિત્વ અસંતુલન',
        'adhd': 'એડીએચડી (ADHD)',
        'autism': 'ઓટિઝમ',
        'eating': 'ખાવાની આદતોની સમસ્યાઓ',
        'main symptoms': 'મુખ્ય લક્ષણો',
        'myths': 'ભ્રમિત ખ્યાલો',
        'routine': 'દૈનિક દિનચર્યા',
        'social': 'સામાજિક પ્રવૃત્તિઓ',
        'food': 'ખોરાકની આદતો'
    }
}

def localize_text(text, lang):
    """
    Filters text to show only the selected language part.
    Supports formats: 'Guj (Eng) | Hindi', 'Guj / Eng / Hindi', 'Eng - Guj'
    Also translates standard headers using PDF_TRANSLATIONS.
    """
    if not text:
        return ""

    # 1. Try Dictionary Lookup (for headers like "Main Symptoms")
    text_lower = text.lower().strip()
    translations = PDF_TRANSLATIONS.get(lang, {})
    
    # Reverse lookup map for English keys to allow mapping "Main Symptoms" -> Localized
    # (Since keys in PDF_TRANSLATIONS['en'] are 'main symptoms', we match keys)
    if text_lower in translations:
        return translations[text_lower]
    
    # Check if it matches a key in the EN dict, then get the Lang dict value
    # e.g. text="Main Symptoms" -> key="main symptoms" -> get from 'gu' dict
    en_keys = PDF_TRANSLATIONS['en']
    found_key = None
    for k, v in en_keys.items():
        if v.lower() == text_lower or k == text_lower:
            found_key = k
            break
            
    if found_key and found_key in translations:
        return translations[found_key]

    # 2. Delimiter Processing
    # Handle delimiters from create_data.py logic: | for paragraphs, / for headers
    if '|' in text:
        parts = [p.strip() for p in text.split('|')]
        # Expected Format: Gujarati (English) | Hindi
        if lang == 'hi':
             return parts[1] if len(parts) > 1 else parts[0]
        
        main_part = parts[0]
        if lang == 'en':
            match = re.search(r'\((.*?)\)', main_part)
            return match.group(1) if match else main_part
        if lang == 'gu':
            return re.sub(r'\(.*?\)', '', main_part).strip()

    if '/' in text:
        parts = [p.strip() for p in text.split('/')]
        # Common formats: G / E / H or G / E or G(E) / H
        if lang == 'hi':
            return parts[-1] if len(parts) > 1 else parts[0]
        if lang == 'en':
            if len(parts) >= 3: return parts[1] # Case: G / E / H
            # Check for G(E) / H
            match = re.search(r'\((.*?)\)', parts[0])
            if match: return match.group(1)
            # Case: G / E -> return E
            return parts[-1] if len(parts) == 2 else parts[0]
        if lang == 'gu':
            return re.sub(r'\(.*?\)', '', parts[0]).strip()

    # Handle "English – Gujarati" format
    if ' – ' in text:
        parts = [p.strip() for p in text.split(' – ')]
        if lang == 'en': return parts[0]
        if lang == 'gu': return parts[1]
        if lang == 'hi': return parts[0] # Fallback to English

    return text

def get_pdf_text(key, lang='en'):
    return PDF_TRANSLATIONS.get(lang, PDF_TRANSLATIONS['en']).get(key, '')

def detect_text_language(text):
    """
    Detect if text is primarily Gujarati, Hindi, or English based on Unicode ranges.
    Returns: 'gu', 'hi', or 'en'
    """
    if not text:
        return 'en'
    
    gujarati_count = sum(1 for char in text if '\u0A80' <= char <= '\u0AFF')  # Gujarati Unicode range
    hindi_count = sum(1 for char in text if '\u0900' <= char <= '\u097F')     # Devanagari (Hindi) Unicode range
    latin_count = sum(1 for char in text if char.isalpha() and ord(char) < 128)  # Basic Latin
    
    total_chars = gujarati_count + hindi_count + latin_count
    if total_chars == 0:
        return 'en'
    
    # Return language with highest character count
    if gujarati_count > hindi_count and gujarati_count > latin_count:
        return 'gu'
    elif hindi_count > gujarati_count and hindi_count > latin_count:
        return 'hi'
    else:
        return 'en'

def extract_advice_from_docx(upload_folder, symptoms, lang='en'):
    """
    Scans the latest uploaded docx file and extracts content ONLY in the selected language.
    Returns a dict: { 'Display Title': [ {'subtitle': 'Subheader', 'content': ['lines']} ] }
    """
    files = [f for f in os.listdir(upload_folder) if f.endswith('.docx')]
    if not files:
        return {"General": [{"subtitle": "Error", "content": ["No advice files found."]}]}
    
    latest_file = max([os.path.join(upload_folder, f) for f in files], key=os.path.getctime)
    
    # Symptom keywords for matching topics (language-agnostic)
    symptom_keywords = {
        'gad': ['anxiety', 'gad', 'ચિંતા', 'चिंता'],
        'depression': ['depression', 'ડિપ્રેશન', 'डिप्रेशन', 'उदासी'],
        'stress': ['stress', 'burnout', 'તણાવ', 'तनाव'],
        'ocd': ['ocd', 'obsessive', 'compulsive', 'બળજબરી', 'ઓસીડી', 'ओसीडी'],
        'ptsd': ['ptsd', 'trauma', 'આઘાત', 'पीटीएसडी'],
        'personality': ['personality', 'વ્યક્તિત્વ', 'व्यक्तित्व'],
        'adhd': ['adhd', 'attention', 'એડીએચડી', 'एडीएचडी'],
        'autism': ['autism', 'ઓટિઝમ', 'ऑटिज्म'],
        'eating': ['eating', 'food disorder', 'ખોરાક', 'ईटिंग']
    }
    
    # Localized display names
    display_names_localized = {
        'en': {
            'gad': 'Generalized Anxiety (GAD)',
            'depression': 'Depression',
            'stress': 'Stress & Burnout',
            'ocd': 'Obsessive-Compulsive (OCD)',
            'ptsd': 'Post-Traumatic Stress (PTSD)',
            'personality': 'Personality Imbalance',
            'adhd': 'ADHD',
            'autism': 'Autism Spectrum',
            'eating': 'Eating Disorders'
        },
        'gu': {
            'gad': 'ચિંતાનો રોગ (GAD)',
            'depression': 'ડિપ્રેશન',
            'stress': 'તણાવ અને થાક',
            'ocd': 'ઓસીડી (OCD)',
            'ptsd': 'પીટીએસડી (PTSD)',
            'personality': 'વ્યક્તિત્વ અસંતુલન',
            'adhd': 'એડીએચડી (ADHD)',
            'autism': 'ઓટિઝમ',
            'eating': 'ખાવાની આદતોની સમસ્યાઓ'
        },
        'hi': {
            'gad': 'सामान्यीकृत चिंता विकार (GAD)',
            'depression': 'अवसाद (Depression)',
            'stress': 'तनाव और बर्नआउट',
            'ocd': 'जुनून-बाध्यकारी विकार (OCD)',
            'ptsd': 'पोस्ट-ट्रॉमैटिक स्ट्रेस (PTSD)',
            'personality': 'व्यक्तित्व असंतुलन',
            'adhd': 'एडीएचडी (ADHD)',
            'autism': 'ऑटिज़्म स्पेक्ट्रम',
            'eating': 'खाने के विकार'
        }
    }
    
    # Section header keywords for each language
    section_headers = {
        'en': {
            'social': ['social activity', 'social activities'],
            'exercise': ['exercise', 'physical activity'],
            'food': ['food habits', 'diet', 'nutrition'],
            'routine': ['routine', 'daily routine', 'mental techniques'],
            'symptoms': ['main symptoms', 'symptoms'],
            'myths': ['myths', 'misconceptions']
        },
        'gu': {
            'social': ['સોશિયલ એક્ટિવિટી', 'સામાજિક પ્રવૃત્તિઓ'],
            'exercise': ['એક્સરસાઈઝ', 'વ્યાયામ'],
            'food': ['ફૂડ હેબિટ્સ', 'ખોરાકની આદતો'],
            'routine': ['રૂટિન', 'દિનચર્યા', 'મનની ટેક્નિક્સ'],
            'symptoms': ['મુખ્ય લક્ષણો', 'લક્ષણો'],
            'myths': ['ભ્રમિત ખ્યાલો', 'ભ્રામક']
        },
        'hi': {
            'social': ['सामाजिक गतिविधियाँ', 'सामाजिक'],
            'exercise': ['व्यायाम', 'एक्सरसाइज'],
            'food': ['भोजन की आदतें', 'आहार'],
            'routine': ['दिनचर्या', 'दैनिक'],
            'symptoms': ['मुख्य लक्षण', 'लक्षण'],
            'myths': ['भ्रामक धारणाएं', 'मिथक']
        }
    }

    advice_structure = {}
    
    try:
        doc = Document(latest_file)
        user_symptoms = [s.strip() for s in symptoms.split(',') if s.strip()]
        
        # Get main symptom categories (remove sub-symptom suffixes)
        main_symptoms = set()
        for sym in user_symptoms:
            main_sym = sym.split('_')[0] if '_' in sym else sym
            if main_sym in symptom_keywords:
                main_symptoms.add(main_sym)
        
        if not main_symptoms:
            return advice_structure

        # Process each symptom
        for symptom_id in main_symptoms:
            keywords = symptom_keywords.get(symptom_id, [symptom_id])
            display_title = display_names_localized.get(lang, {}).get(symptom_id, symptom_id.title())
            
            advice_structure[display_title] = []
            
            capturing = False
            current_lang_section = None
            current_subtopic = None
            seen_sections = {}  # Track sections to avoid duplicates
            
            for para in doc.paragraphs:
                text = para.text.strip()
                if not text:
                    continue
                
                text_lower = text.lower()
                detected_lang = detect_text_language(text)
                
                # Check if this is a main topic header (matches our symptom)
                is_topic_match = any(keyword in text_lower for keyword in keywords)
                
                # Check if this matches a DIFFERENT disorder (to stop capturing)
                is_other_disorder = False
                if len(text) < 150:
                    for other_id, other_keywords in symptom_keywords.items():
                        if other_id != symptom_id:
                            if any(kw in text_lower for kw in other_keywords):
                                is_other_disorder = True
                                break
                
                if is_topic_match and len(text) < 150:
                    # Starting a new topic section
                    if detected_lang == lang:
                        # This is the language section we want
                        capturing = True
                        current_lang_section = lang
                        current_subtopic = None
                        seen_sections = {}
                    elif capturing and detected_lang != lang:
                        # We've moved to a different language section of the same topic - stop
                        if current_subtopic and current_subtopic.get('content'):
                            advice_structure[display_title].append(current_subtopic)
                        capturing = False
                        current_lang_section = None
                    continue
                
                # Stop if we hit a different disorder topic
                if capturing and is_other_disorder:
                    if current_subtopic and current_subtopic.get('content'):
                        advice_structure[display_title].append(current_subtopic)
                    capturing = False
                    current_lang_section = None
                    continue
                
                # If we're capturing and in the right language section
                if capturing and current_lang_section == lang:
                    # Check if this is a section header (Social Activity, Food Habits, etc.)
                    is_section_header = False
                    section_key_found = None
                    
                    for section_key, header_variants in section_headers.get(lang, {}).items():
                        if any(header.lower() in text_lower for header in header_variants):
                            is_section_header = True
                            section_key_found = section_key
                            break
                    
                    # Also check for numbered headers like "1. Social Activity"
                    if not is_section_header and len(text) < 100:
                        # Remove numbering and check again
                        clean_header = re.sub(r'^\d+\.\s*', '', text).strip()
                        for section_key, header_variants in section_headers.get(lang, {}).items():
                            if any(header.lower() in clean_header.lower() for header in header_variants):
                                is_section_header = True
                                section_key_found = section_key
                                break
                    
                    if is_section_header and section_key_found:
                        # Save previous subtopic if exists
                        if current_subtopic and current_subtopic.get('content'):
                            # Check if we already have this section
                            standardized_name = get_pdf_text(section_key_found, lang) or text
                            if standardized_name in seen_sections:
                                # Merge content into existing section
                                seen_sections[standardized_name]['content'].extend(current_subtopic['content'])
                            else:
                                advice_structure[display_title].append(current_subtopic)
                                seen_sections[current_subtopic['subtitle']] = current_subtopic
                        
                        # Use standardized section name from PDF_TRANSLATIONS
                        standardized_name = get_pdf_text(section_key_found, lang) or text
                        
                        # Check if this section already exists
                        if standardized_name in seen_sections:
                            # Reuse existing section
                            current_subtopic = seen_sections[standardized_name]
                        else:
                            # Start new subtopic
                            current_subtopic = {'subtitle': standardized_name, 'content': []}
                            seen_sections[standardized_name] = current_subtopic
                    else:
                        # This is content - add to current subtopic
                        if detected_lang == lang:  # Only add if it's in the target language
                            clean_text = text.lstrip('-•').strip()
                            if clean_text and len(clean_text) > 3:  # Avoid very short fragments
                                if current_subtopic is None:
                                    general_title = get_pdf_text('no_data', lang) if lang != 'en' else 'General'
                                    current_subtopic = {'subtitle': general_title, 'content': []}
                                    seen_sections[general_title] = current_subtopic
                                    seen_sections[general_title] = current_subtopic
                                
                                # Deduplicate content at insertion
                                if clean_text not in current_subtopic['content']:
                                    current_subtopic['content'].append(clean_text)
            
            # Save last subtopic if not already saved
            if current_subtopic and current_subtopic.get('content'):
                if current_subtopic not in advice_structure[display_title]:
                    advice_structure[display_title].append(current_subtopic)
            
            # Deduplicate - keep only unique sections and unique content within sections
            unique_sections = []
            seen_titles = set()
            for section in advice_structure[display_title]:
                if section['subtitle'] not in seen_titles:
                    # Deduplicate content lines while preserving order
                    seen_lines = set()
                    unique_content = []
                    for line in section['content']:
                        if line not in seen_lines:
                            unique_content.append(line)
                            seen_lines.add(line)
                    section['content'] = unique_content
                    
                    if unique_content: # Only keep if there is content
                        unique_sections.append(section)
                        seen_titles.add(section['subtitle'])
            
            advice_structure[display_title] = unique_sections
            
            # Remove empty entries
            if not advice_structure[display_title]:
                del advice_structure[display_title]

    except Exception as e:
        return {"Error": [{"subtitle": "Error", "content": [f"Error reading doc: {str(e)}"]}]}
        
    return advice_structure


class PrajnaPathPDF(FPDF):
    def __init__(self, language='en'):
        super().__init__()
        self.language = language
        self.set_auto_page_break(auto=True, margin=15)
        
        # Add fonts
        if os.path.exists(FONT_PATH_NIRMALA):
            self.add_font("Nirmala", style="", fname=FONT_PATH_NIRMALA)
            self.add_font("Nirmala", style="B", fname=FONT_PATH_NIRMALA)
        else:
            # Fallback (unicode won't work well)
            self.add_font("Arial", "", "arial.ttf", uni=True) 

    def header(self):
        # Professional Header with Peacock Teal Background
        self.set_fill_color(13, 92, 99) # #0d5c63
        self.rect(0, 0, 210, 40, 'F')
        
        self.set_font('Nirmala', 'B', 24)
        self.set_text_color(255, 255, 255) # White
        
        # Localized Title
        title = get_pdf_text('title', self.language)
        self.cell(0, 20, title, 0, 1, 'C')
        
        self.set_font('Nirmala', '', 12)
        self.set_text_color(240, 240, 240)
        subtitle = get_pdf_text('subtitle', self.language)
        self.cell(0, 5, subtitle, 0, 1, 'C')
        
        self.ln(20)

    def footer(self):
        self.set_y(-15)
        self.set_font('Nirmala', '', 8)
        self.set_text_color(128, 128, 128)
        
        disclaimer = get_pdf_text('disclaimer', self.language)
        page_txt = get_pdf_text('page', self.language)
        
        self.cell(0, 5, disclaimer, 0, 1, 'C')
        self.cell(0, 5, f'{page_txt} {self.page_no()}', 0, 0, 'C')

def create_pdf_report(user_session, advice_data):
    # Detect language from user_session (passed as string code 'en', 'hi', 'gu')
    lang = getattr(user_session, 'language', 'en')
    
    pdf = PrajnaPathPDF(language=lang)
    pdf.add_page()
    
    # --- User Details Section ---
    pdf.set_text_color(0, 0, 0)
    pdf.set_font("Nirmala", "B", 14)
    # Color: Peacock Teal
    pdf.set_text_color(13, 92, 99)
    
    # Ensure we start BELOW the header (header height is 40)
    pdf.set_y(45)
    
    # "Personal Profile" or similar if needed, but 'title' is already in header
    # pdf.cell(0, 10,  get_pdf_text('title', lang), 0, 1) 
    
    pdf.set_text_color(50, 50, 50)
    pdf.set_font("Nirmala", "", 12)
    
    # Draw a box for details
    pdf.set_draw_color(200, 200, 200)
    pdf.rect(10, 50, 190, 35)
    
    start_y = 55
    pdf.set_xy(15, start_y)
    
    # Labels always in English
    labels = {
        'name': 'Name',
        'age': 'Age',
        'prof': 'Profession',
        'date': 'Date'
    }
    
    # Safe text handling
    def safe_txt(text):
        try:
            return str(text) # Ensure string
        except:
            return ""

    pdf.cell(90, 8, f"{labels['name']}: {safe_txt(user_session.name)}", 0, 0)
    pdf.cell(90, 8, f"{labels['age']}: {safe_txt(user_session.age)}", 0, 1)
    
    pdf.set_x(15)
    pdf.cell(90, 8, f"{labels['prof']}: {safe_txt(user_session.profession)}", 0, 0)
    pdf.cell(90, 8, f"{labels['date']}: {datetime.now().strftime('%Y-%m-%d')}", 0, 1)
    
    pdf.ln(15)
    
    for symptom, subtopics in advice_data.items():
        # The symptom title is already localized in extract_advice_from_docx
        display_symptom = symptom
        
        # Section Header
        pdf.set_fill_color(255, 153, 51) # Saffron accent
        pdf.rect(pdf.get_x(), pdf.get_y(), 3, 10, 'F') # Color bar on left
        
        pdf.set_x(pdf.get_x() + 5)
        pdf.set_font("Nirmala", 'B', 16)
        pdf.set_text_color(13, 92, 99) # Teal
        
        try:
            pdf.cell(0, 10, f"{display_symptom}", 0, 1)
        except:
            safe = display_symptom.encode('latin-1', 'replace').decode('latin-1')
            pdf.cell(0, 10, f"{safe}", 0, 1)
            
        pdf.ln(2)
        pdf.set_draw_color(13, 92, 99)
        pdf.line(10, pdf.get_y(), 200, pdf.get_y()) # Underline
        pdf.ln(5)

        if not subtopics:
             continue

        for sub in subtopics:
            # Subtitle (already localized)
            if sub.get('subtitle'):
                display_subtitle = sub['subtitle']
                
                # Check if this is a special section to highlight
                special_keywords = ['Routine', 'Social', 'Food', 'દિનચર્યા', 'પ્રવૃત્તિઓ', 'ખોરાક', 'दिनचर्या', 'गतिविधियाँ', 'भोजन']
                is_special = any(k.lower() in display_subtitle.lower() for k in special_keywords)

                if is_special:
                    pdf.set_fill_color(255, 245, 230) # Very light saffron background
                    pdf.set_font("Nirmala", 'B', 13)
                    pdf.set_text_color(204, 85, 0) # Saffron
                    pdf.ln(3)
                    pdf.cell(0, 10, f"  {display_subtitle}", 0, 1, 'L', fill=True)
                else:
                    pdf.set_font("Nirmala", 'B', 13)
                    pdf.set_text_color(204, 85, 0) # Saffron-ish
                    pdf.ln(3)
                    try:
                        pdf.cell(0, 8, f"{display_subtitle}", 0, 1)
                    except:
                        pdf.set_font("Helvetica", 'B', 12) # Fallback
                        pdf.cell(0, 8, f"{display_subtitle}", 0, 1)

            # Bullet content
            pdf.set_font("Nirmala", size=11)
            pdf.set_text_color(20, 20, 20)
            
            for item in sub['content']:
                # The item is already localized during extraction
                clean_item = item
                
                pdf.set_x(15) # Indent
                # Bullet point
                pdf.cell(5, 6, chr(149), 0, 0) 
                
                try:
                    pdf.multi_cell(0, 6, f"{clean_item}")
                except Exception:
                    # Robust fallback for font errors
                    current = pdf.font_family
                    pdf.set_font("Helvetica", size=11)
                    safe_item = clean_item.encode('latin-1', 'replace').decode('latin-1')
                    pdf.multi_cell(0, 6, f"{safe_item}")
                    pdf.set_font(current, size=11)
                    
            pdf.ln(3)
            
        pdf.ln(8) # Space between topics

    filename = f"report_{user_session.id}.pdf"
    path = os.path.join("static", filename)
    pdf.output(path)
    return path

